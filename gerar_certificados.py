"""Prepara rascunhos de certificados no Outlook Desktop, sem enviar mensagens.

Uso:
    python gerar_certificados.py "C:\Certificados" --dry-run
    python gerar_certificados.py "C:\Certificados" --relatorio-csv resultado.csv

A pasta deve conter planilha_envio_certificados.xlsx (aba Envio) e os PDFs.
"""

import argparse
import csv
from contextlib import contextmanager
from datetime import date, datetime
from html import escape
from pathlib import Path, PureWindowsPath
import re

from openpyxl import load_workbook


PLANILHA_NOME = "planilha_envio_certificados.xlsx"
ABA = "Envio"
ASSUNTO = "Seu certificado do {Nome do Curso} já está disponível! 🎓"
CORPO = """
<p>Olá, {Nome do Destinatário},</p>
<p>Esperamos que este e-mail o(a) encontre bem.</p>
<p>Em nome do Centro de Treinamento Médico, gostaríamos de agradecer imensamente pela sua participação
 e dedicação no {Nome do Curso}, realizado em {Data do Curso}.</p>
<p>Sua presença e colaboração como {Tipo Participação} foram fundamentais para o sucesso e a excelência
 do nosso evento.</p>
<p>Segue em anexo o seu certificado.</p>
<p>Atenciosamente,<br>Centro de Treinamento Médico</p>
"""
CAMPOS_OBRIGATORIOS = {
    "Nome do Destinatário", "Email", "Anexo", "Nome do Curso",
    "Data do Curso", "Tipo Participação",
}
PADRAO_EMAIL = re.compile(r"^[^\s@;,]+@[^\s@;,]+\.[^\s@;,]+$")
PADRAO_TAG = re.compile(r"\{([^{}]+)\}")


def texto_celula(valor):
    """Converte células vazias e datas do Excel para texto legível."""
    if valor is None:
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip()


def carregar_linhas(caminho_planilha):
    """Lê a aba Envio e impede planilhas sem as colunas essenciais."""
    pasta = load_workbook(caminho_planilha, read_only=True, data_only=True)
    try:
        if ABA not in pasta.sheetnames:
            raise ValueError(f"A planilha precisa conter a aba '{ABA}'.")
        linhas = pasta[ABA].iter_rows(values_only=True)
        cabecalho = next(linhas, None)
        if not cabecalho:
            raise ValueError("A aba 'Envio' está vazia.")
        colunas = [texto_celula(coluna) for coluna in cabecalho]
        ausentes = CAMPOS_OBRIGATORIOS - set(colunas)
        if ausentes:
            raise ValueError("Colunas ausentes: " + ", ".join(sorted(ausentes)))
        if len(set(colunas)) != len(colunas):
            raise ValueError("Há nomes de coluna repetidos na aba 'Envio'.")
        return [
            (numero, dict(zip(colunas, valores)))
            for numero, valores in enumerate(linhas, start=2)
            if any(texto_celula(valor) for valor in valores)
        ]
    finally:
        pasta.close()


def localizar_pdf(pasta_anexos, nome_arquivo):
    """Aceita somente PDFs existentes dentro da pasta selecionada."""
    nome = texto_celula(nome_arquivo)
    if not nome:
        raise ValueError("Nome do anexo não informado.")
    solicitado = Path(nome)
    if solicitado.is_absolute() or PureWindowsPath(nome).is_absolute():
        raise ValueError("O anexo deve ter caminho relativo à pasta de certificados.")
    if ".." in solicitado.parts or ".." in PureWindowsPath(nome).parts:
        raise ValueError("O caminho do anexo não pode sair da pasta de certificados.")
    if solicitado.suffix.lower() != ".pdf":
        raise ValueError("O certificado deve estar em PDF.")
    raiz = Path(pasta_anexos).resolve()
    arquivo = (raiz / solicitado).resolve()
    try:
        arquivo.relative_to(raiz)
    except ValueError as erro:
        raise ValueError("O caminho do anexo não pode sair da pasta de certificados.") from erro
    if not arquivo.is_file():
        raise ValueError(f"PDF não encontrado: {nome}")
    return arquivo


def substituir_tags(modelo, dados, html=False):
    """Substitui as tags uma vez, sem interpretar outras tags dentro dos dados."""
    def substituir(correspondencia):
        valor = texto_celula(dados.get(correspondencia.group(1)))
        return escape(valor, quote=True) if html else valor
    return PADRAO_TAG.sub(substituir, modelo)


def preparar_lote(pasta_anexos, linhas):
    """Consolida os itens válidos antes de abrir qualquer janela do Outlook."""
    aprovados = []
    relatorio = []
    duplicados = set()
    for numero, dados in linhas:
        nome = texto_celula(dados.get("Nome do Destinatário"))
        email = texto_celula(dados.get("Email"))
        anexo = texto_celula(dados.get("Anexo"))
        problemas = []
        if not PADRAO_EMAIL.fullmatch(email):
            problemas.append("E-mail inválido ou vazio.")
        for campo in ("Nome do Destinatário", "Nome do Curso", "Data do Curso", "Tipo Participação"):
            if not texto_celula(dados.get(campo)):
                problemas.append(f"Campo obrigatório vazio: {campo}.")
        arquivo = None
        try:
            arquivo = localizar_pdf(pasta_anexos, anexo)
        except ValueError as erro:
            problemas.append(str(erro))

        if not problemas:
            chave = (email.casefold(), str(arquivo).casefold())
            if chave in duplicados:
                problemas.append("Registro duplicado para o mesmo e-mail e PDF.")
            else:
                duplicados.add(chave)

        registro = {
            "linha": numero,
            "nome": nome,
            "email": email,
            "anexo": anexo,
            "status": "Ignorado" if problemas else "Pronto",
            "detalhe": " ".join(problemas),
        }
        relatorio.append(registro)
        if not problemas:
            aprovados.append({
                "registro": registro,
                "arquivo": arquivo,
                "assunto": substituir_tags(ASSUNTO, dados),
                "corpo": substituir_tags(CORPO, dados, html=True),
            })
    return aprovados, relatorio


@contextmanager
def conectar_outlook():
    """Garante liberação do contexto COM mesmo em caso de erro."""
    import pythoncom
    import win32com.client as win32

    pythoncom.CoInitialize()
    try:
        try:
            outlook = win32.GetActiveObject("Outlook.Application")
        except Exception:
            outlook = win32.Dispatch("Outlook.Application")
        yield outlook
    finally:
        pythoncom.CoUninitialize()


def criar_rascunho(outlook, item):
    """Cria um único rascunho, preservando a assinatura do Outlook."""
    mensagem = outlook.CreateItem(0)
    try:
        mensagem.Display()
        mensagem.To = item["registro"]["email"]
        mensagem.Subject = item["assunto"]
        mensagem.HTMLBody = (
            "<div style='font-family: Calibri; font-size: 11pt;'>"
            + item["corpo"] + "</div><br>" + mensagem.HTMLBody
        )
        mensagem.Attachments.Add(str(item["arquivo"]))
        mensagem.Save()
        mensagem.Close(0)
    except Exception:
        try:
            mensagem.Close(1)  # descarta a janela parcialmente preenchida
        except Exception:
            pass
        raise


def gravar_relatorio(caminho, linhas):
    """Grava um CSV opcional com status por linha (contém dados pessoais)."""
    with open(caminho, "w", encoding="utf-8-sig", newline="") as destino:
        escritor = csv.DictWriter(
            destino, fieldnames=["linha", "nome", "email", "anexo", "status", "detalhe"],
            delimiter=";",
        )
        escritor.writeheader()
        escritor.writerows(linhas)


def main(argv=None):
    parser = argparse.ArgumentParser(description="Gera rascunhos de certificados no Outlook Desktop.")
    parser.add_argument("pasta", help="Pasta que contém a planilha e os PDFs")
    parser.add_argument("--dry-run", action="store_true", help="Confere o lote sem abrir o Outlook")
    parser.add_argument("--relatorio-csv", help="Caminho opcional para o relatório por linha")
    args = parser.parse_args(argv)

    pasta = Path(args.pasta).expanduser()
    planilha = pasta / PLANILHA_NOME
    if not pasta.is_dir() or not planilha.is_file():
        parser.error(f"Verifique a pasta e a planilha '{PLANILHA_NOME}'.")

    try:
        aprovados, relatorio = preparar_lote(pasta, carregar_linhas(planilha))
    except (OSError, ValueError) as erro:
        parser.error(str(erro))

    for linha in relatorio:
        print(f"Linha {linha['linha']}: {linha['status']} - {linha['email']} {linha['detalhe']}")

    if args.dry_run:
        print(f"Simulação: {len(aprovados)} rascunhos aptos; nenhum rascunho foi criado.")
    elif aprovados:
        try:
            with conectar_outlook() as outlook:
                for item in aprovados:
                    try:
                        criar_rascunho(outlook, item)
                        item["registro"]["status"] = "Criado"
                    except Exception as erro:
                        item["registro"]["status"] = "Erro"
                        item["registro"]["detalhe"] = str(erro)
                        print(f"Falha ao criar rascunho da linha {item['registro']['linha']}: {erro}")
        except Exception as erro:
            print(f"Não foi possível conectar ao Outlook Desktop: {erro}")
            for item in aprovados:
                if item["registro"]["status"] == "Pronto":
                    item["registro"]["status"] = "Erro"
                    item["registro"]["detalhe"] = "Conexão indisponível com Outlook Desktop."
    else:
        print("Nenhum registro válido; nenhum rascunho foi criado.")

    if args.relatorio_csv:
        try:
            gravar_relatorio(args.relatorio_csv, relatorio)
            print(f"Relatório salvo: {args.relatorio_csv}")
        except OSError as erro:
            print(f"Não foi possível gravar o relatório: {erro}")
            return 1

    if not args.dry_run:
        criados = sum(linha["status"] == "Criado" for linha in relatorio)
        print(f"Rascunhos criados: {criados}; registros ignorados/com erro: {len(relatorio) - criados}.")
    return 0 if relatorio and all(
        linha["status"] == ("Pronto" if args.dry_run else "Criado") for linha in relatorio
    ) else 1


if __name__ == "__main__":
    raise SystemExit(main())
